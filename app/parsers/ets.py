from __future__ import annotations

import logging
import re
from datetime import date, datetime
from io import BytesIO
from typing import Optional
from urllib.parse import urljoin

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from app import config
from app.parsers.http_client import HttpClient
from app.parsers.normalize import (
    NormalizedTrade,
    clean_text,
    parse_date,
    parse_number,
)

log = logging.getLogger(__name__)

MONTH_SLUG_RE = re.compile(r"20\d{2}")
XLSX_RE = re.compile(r"\.xlsx($|\?)", re.I)
DATE_IN_TITLE_RE = re.compile(r"(\d{2}\.\d{2}\.\d{4})")


class EtsParser:
    exchange_code = "ETS"

    def __init__(self, client: Optional[HttpClient] = None) -> None:
        self.client = client or HttpClient()

    def parse(
        self,
        *,
        mode: str = "incremental",
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
    ) -> list[NormalizedTrade]:
        date_from, date_to = self._resolve_range(mode, date_from, date_to)
        log.info("ETS parse mode=%s range=%s..%s", mode, date_from, date_to)
        files = self.discover_xlsx_files(date_from, date_to)
        log.info("ETS found %s xlsx files", len(files))
        trades: list[NormalizedTrade] = []
        for title, url, trade_date in files:
            try:
                data = self.client.get_bytes(url)
                rows = self.parse_xlsx(data, trade_date=trade_date, source_url=url)
                trades.extend(rows)
                log.info("ETS %s -> %s rows", title[:60], len(rows))
            except Exception as exc:  # noqa: BLE001
                log.exception("ETS failed file %s: %s", url, exc)
        return trades

    def _resolve_range(
        self,
        mode: str,
        date_from: Optional[date],
        date_to: Optional[date],
    ) -> tuple[Optional[date], Optional[date]]:
        if mode == "full" and date_from is None and date_to is None:
            return None, None
        if date_to is None:
            date_to = date.today()
        if date_from is None:
            # последние N месяцев
            months = config.ETS_INCREMENTAL_MONTHS
            y, m = date_to.year, date_to.month - months
            while m <= 0:
                m += 12
                y -= 1
            date_from = date(y, m, 1)
        return date_from, date_to

    def discover_xlsx_files(
        self,
        date_from: Optional[date],
        date_to: Optional[date],
    ) -> list[tuple[str, str, date]]:
        catalog_html = self.client.get_text(config.ETS_CATALOG_URL)
        year_pages = self._extract_year_links(catalog_html, config.ETS_CATALOG_URL)
        results: list[tuple[str, str, date]] = []
        for ytitle, yurl in year_pages:
            year_num = self._year_from_text(ytitle + yurl)
            if date_from and year_num and year_num < date_from.year:
                continue
            if date_to and year_num and year_num > date_to.year:
                continue
            try:
                yhtml = self.client.get_text(yurl)
            except Exception as exc:  # noqa: BLE001
                log.error("ETS year page fail %s: %s", yurl, exc)
                continue
            month_pages = self._extract_month_links(yhtml, yurl)
            for mtitle, murl in month_pages:
                try:
                    mhtml = self.client.get_text(murl)
                except Exception as exc:  # noqa: BLE001
                    log.error("ETS month page fail %s: %s", murl, exc)
                    continue
                for title, furl, tdate in self._extract_file_links(mhtml, murl):
                    if date_from and tdate < date_from:
                        continue
                    if date_to and tdate > date_to:
                        continue
                    results.append((title, furl, tdate))
        # уникальные URL
        seen: set[str] = set()
        unique: list[tuple[str, str, date]] = []
        for item in results:
            if item[1] in seen:
                continue
            seen.add(item[1])
            unique.append(item)
        return unique

    def _extract_year_links(self, html: str, base: str) -> list[tuple[str, str]]:
        soup = BeautifulSoup(html, "lxml")
        out: list[tuple[str, str]] = []
        seen: set[str] = set()
        for a in soup.find_all("a", href=True):
            href = a["href"]
            text = a.get_text(" ", strip=True)
            full = urljoin(base, href)
            if "trading-results" not in full:
                continue
            if full.rstrip("/") == config.ETS_CATALOG_URL.rstrip("/"):
                continue
            if not re.search(r"20\d{2}", text + full):
                continue
            # только прямые дочерние годовые разделы
            path = full.replace(config.ETS_BASE_URL, "")
            depth = path.strip("/").count("/")
            if depth > 3:
                continue
            if full in seen:
                continue
            # отсекаем языковые зеркала
            if "/kz/" in full or "/en/" in full:
                continue
            seen.add(full)
            out.append((text, full))
        return out

    def _extract_month_links(self, html: str, base: str) -> list[tuple[str, str]]:
        soup = BeautifulSoup(html, "lxml")
        out: list[tuple[str, str]] = []
        seen: set[str] = set()
        for a in soup.find_all("a", href=True):
            href = a["href"]
            text = a.get_text(" ", strip=True)
            full = urljoin(base, href)
            if full.rstrip("/") == base.rstrip("/"):
                continue
            if "trading-results" not in full:
                continue
            if "/kz/" in full or "/en/" in full:
                continue
            if XLSX_RE.search(full):
                continue
            # месячная страница глубже годовой
            if full.count("/") <= base.count("/"):
                continue
            if full in seen:
                continue
            seen.add(full)
            out.append((text, full))
        return out

    def _extract_file_links(
        self, html: str, base: str
    ) -> list[tuple[str, str, date]]:
        soup = BeautifulSoup(html, "lxml")
        out: list[tuple[str, str, date]] = []
        for a in soup.find_all("a", href=True):
            href = a["href"]
            if not (XLSX_RE.search(href) or (".xlsx" in href.lower())):
                continue
            full = urljoin(base, href)
            # заголовок документа — из ближайшего контейнера
            parent = a.find_parent(["div", "tr", "li", "td", "article"])
            title = parent.get_text(" ", strip=True) if parent else a.get_text(strip=True)
            m = DATE_IN_TITLE_RE.search(title)
            if not m:
                # иногда дата только рядом
                prev = a.find_previous(string=DATE_IN_TITLE_RE)
                if prev:
                    m = DATE_IN_TITLE_RE.search(str(prev))
            if not m:
                log.warning("ETS skip file without date: %s", full)
                continue
            tdate = parse_date(m.group(1))
            if not tdate:
                continue
            out.append((title, full, tdate))
        return out

    @staticmethod
    def _year_from_text(text: str) -> Optional[int]:
        m = re.search(r"(20\d{2})", text)
        return int(m.group(1)) if m else None

    def parse_xlsx(
        self,
        data: bytes,
        *,
        trade_date: date,
        source_url: str,
    ) -> list[NormalizedTrade]:
        wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
        trades: list[NormalizedTrade] = []
        for sheet in wb.worksheets:
            # дата из имени листа, если есть
            sheet_date = parse_date(sheet.title) or trade_date
            header_row_idx = None
            headers: dict[int, str] = {}
            for i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = list(row)
                joined = " ".join(str(v) for v in values if v is not None)
                if "Наименование покупателя" in joined or "Сатып алушының атауы" in joined:
                    header_row_idx = i
                    for idx, cell in enumerate(values):
                        if cell is not None:
                            headers[idx] = str(cell)
                    continue
                if header_row_idx is None:
                    # иногда дата в заголовке листа
                    dm = DATE_IN_TITLE_RE.search(joined)
                    if dm:
                        sheet_date = parse_date(dm.group(1)) or sheet_date
                    continue
                mapped = self._map_row(headers, values)
                if not mapped:
                    continue
                if mapped.get("_skip"):
                    continue
                nt = NormalizedTrade(
                    exchange_code=self.exchange_code,
                    trade_date=sheet_date,
                    product_name=mapped.get("product_name"),
                    product_code=mapped.get("product_code"),
                    tnved_code=mapped.get("tnved"),
                    seller_name=mapped.get("seller_name"),
                    seller_bin=mapped.get("seller_bin"),
                    buyer_name=mapped.get("buyer_name"),
                    buyer_bin=mapped.get("buyer_bin"),
                    seller_broker=mapped.get("seller_broker"),
                    buyer_broker=mapped.get("buyer_broker"),
                    price=mapped.get("price"),
                    amount=mapped.get("amount"),
                    deals_count=mapped.get("deals_count"),
                    source_url=source_url,
                    price_vat=None,
                )
                trades.append(nt.finalize())
        wb.close()
        return trades

    def _map_row(self, headers: dict[int, str], values: list) -> Optional[dict]:
        def cell(idx: int):
            return values[idx] if idx < len(values) else None

        # найти индексы по ключевым словам заголовка
        def find(*keys: str) -> Optional[int]:
            for idx, h in headers.items():
                hl = h.lower()
                if all(k.lower() in hl for k in keys):
                    return idx
            for idx, h in headers.items():
                hl = h.lower()
                if any(k.lower() in hl for k in keys):
                    return idx
            return None

        # пустая / итого
        product_idx = find("наименование товара", "тауардың атауы") or find("товара")
        product_val = clean_text(cell(product_idx)) if product_idx is not None else None
        if product_val and product_val.casefold() in {"итого", "жиыны", "total"}:
            return {"_skip": True}
        # если почти всё пусто
        non_empty = [v for v in values if v not in (None, "")]
        if len(non_empty) < 3:
            return None

        buyer_name_i = find("наименование покупателя") or find("сатып алушының атауы")
        buyer_bin_i = find("бин покупателя") or find("сатып алушының бсн")
        buyer_broker_i = find("брокера покупателя") or find("сатып алушы брокері")
        seller_name_i = find("наименование продавца") or find("сатушының атауы")
        seller_bin_i = find("бин продавца") or find("сатушының бсн")
        seller_broker_i = find("брокера продавца") or find("сатушы брокері")
        tnved_i = find("тн вэд") or find("сэқ тн") or find("еаэо")
        code_i = find("код товара") or find("тауар коды")
        deals_i = find("количество") or find("мәмілелер саны")
        price_i = find("средневзвешанная") or find("средневзвешенная") or find("орташа")
        amount_i = find("общая сумма") or find("жалпы сомасы")

        buyer_name = clean_text(cell(buyer_name_i)) if buyer_name_i is not None else None
        if not buyer_name and not product_val:
            return None

        deals = parse_number(cell(deals_i)) if deals_i is not None else None
        return {
            "buyer_name": buyer_name,
            "buyer_bin": clean_text(cell(buyer_bin_i)) if buyer_bin_i is not None else None,
            "buyer_broker": clean_text(cell(buyer_broker_i)) if buyer_broker_i is not None else None,
            "seller_name": clean_text(cell(seller_name_i)) if seller_name_i is not None else None,
            "seller_bin": clean_text(cell(seller_bin_i)) if seller_bin_i is not None else None,
            "seller_broker": clean_text(cell(seller_broker_i)) if seller_broker_i is not None else None,
            "product_name": product_val,
            "tnved": clean_text(cell(tnved_i)) if tnved_i is not None else None,
            "product_code": clean_text(cell(code_i)) if code_i is not None else None,
            "deals_count": int(deals) if deals is not None else None,
            "price": parse_number(cell(price_i)) if price_i is not None else None,
            "amount": parse_number(cell(amount_i)) if amount_i is not None else None,
        }
